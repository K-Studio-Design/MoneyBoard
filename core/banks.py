"""Banche preconfigurate per l'import degli estratti (CSV / Excel / formati standard).

Ogni banca ha un «profilo»: come si esporta il file dal sito o dall'app e quali colonne usa (nomi possibili
dell'intestazione, ordine giorno/mese, separatore decimale, segno degli importi delle carte). Le banche cambiano
spesso i loro file: per questo, oltre ai nomi propri di ogni banca, si cercano sempre anche i nomi più comuni in
italiano, inglese, tedesco, francese, spagnolo e olandese, e i formati standard (OFX/QFX, QIF, CAMT.053, MT940)
sono accettati per qualsiasi banca.

Fonti dei formati: documentazione delle banche e importatori open source che leggono gli stessi file
(firefly-iii/import-configurations, bank2ynab, csv2ofx, ofxstatement-*, beancount-*).

Fineco, Revolut e Chase hanno un lettore dedicato (core/importers.py, core/chase.py).
"""
from __future__ import annotations

import csv
import io
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta

import pandas as pd

from core import i18n

# ---------------------------------------------------------------- nomi delle colonne (sinonimi)
AUTO: dict[str, list[str]] = {
    "date": [
        "data operazione", "data contabile", "data registrazione", "data movimento", "data", "data op.",
        "transaction date", "trans. date", "trans date", "date", "posted date", "posting date", "booking date",
        "date posted", "completed date", "buchungstag", "buchungsdatum", "buchung", "datum", "transactiedatum",
        "boekdatum", "date de l'operation", "date de l'opération", "date operation", "date opération",
        "date d'operation", "date d'opération", "dateop", "fecha operacion", "fecha operación", "fecha",
        "f. operacion", "f. operación", "fecha de operacion", "fecha de operación", "abschlussdatum",
    ],
    "date2": [
        "data valuta", "value date", "wertstellung", "wertstellungsdatum", "valutadatum", "wert", "valuta",
        "date de valeur", "date valeur", "dateval", "fecha valor", "f. valor", "f.valor", "rentedatum", "post date",
    ],
    "amount": [
        "importo", "importo (eur)", "importo eur", "importo euro", "importo in euro", "importo €", "euro",
        "amount", "amount (eur)", "amount (gbp)", "amount (usd)", "amount (chf)", "transaction amount", "value",
        "betrag", "betrag (eur)", "betrag (€)", "umsatz", "umsatz in eur", "buchungsbetrag", "einzelbetrag",
        "montant", "montant (eur)", "montant de l'operation", "montant de l'opération", "montant operation",
        "montant opération", "importe", "importe (€)", "importe (eur)", "importe eur", "bedrag", "bedrag (eur)",
        "importo (*", "amount (*", "betrag (*", "montant (*", "importe (*", "bedrag (*", "umsatz in *",
    ],
    "debit": [
        "uscite", "addebiti", "addebito", "dare", "importo dare", "debit", "debit amount", "debits", "paid out",
        "money out", "withdrawal", "withdrawals", "withdrawal amount", "soll", "lastschrift", "belastung",
        "debit euros", "débit euros", "debit", "débit", "cargo", "cargos", "debe", "lastschrift in *",
        "belastung *", "débit *", "debit *", "soll (*",
    ],
    "credit": [
        "entrate", "accrediti", "accredito", "avere", "importo avere", "credit", "credit amount", "credits",
        "paid in", "money in", "deposit", "deposits", "deposit amount", "haben", "gutschrift", "credit euros",
        "crédit euros", "crédit", "abono", "abonos", "haber", "ingresos", "gutschrift in *", "gutschrift *",
        "crédit *", "credit *", "haben (*",
    ],
    "dc": [
        "af bij", "af/bij", "debit/credit", "credit/debit", "creditdebet", "cr/dr", "dr/cr", "soll/haben", "s/h",
        "segno", "d/a", "dare/avere", "transaction type", "debit or credit", "credit or debit", "c/d", "d/c",
        "tipo movimento",
    ],
    "ccy": [
        "divisa", "currency", "ccy", "ccy.", "währung", "waehrung", "whg", "devise", "moneda", "munt",
        "buchungswährung", "currency code", "valuta conto",
    ],
    "status": ["stato", "status", "state", "estado", "statut", "contabilizzazione"],
    "category": [
        "categoria", "category", "kategorie", "catégorie", "categorie", "categoría", "spending category",
        "moneymap", "category name", "sottocategoria",
    ],
    "payee": [
        "nome", "beneficiario", "controparte", "ordinante/beneficiario", "payee", "name", "merchant",
        "merchant name", "counter party", "counterparty", "partner name", "auftraggeber/empfänger",
        "auftraggeber / begünstigter", "begünstigter / auftraggeber", "beguenstigter/zahlungspflichtiger",
        "begünstigter/zahlungspflichtiger", "auftraggeber/beguenstigter", "auftraggeber/begünstigter",
        "name zahlungsbeteiligter", "empfänger", "naam tegenpartij", "naam / omschrijving", "name / description",
        "naam", "bénéficiaire", "tiers", "nombre", "ordenante/beneficiario", "payee name", "partnername",
        "partner", "empfänger/zahlungspflichtiger", "auftraggeber/zahlungsempfänger", "zahlungsempfänger",
    ],
    "purpose": [
        "descrizione completa", "descrizione estesa", "descrizione operazioni", "descrizione operazione",
        "descrizione", "dettagli", "causale", "operazione", "description", "transaction description",
        "details", "memo", "reference", "payment reference", "narrative", "particulars", "transaction details",
        "verwendungszweck", "buchungstext", "vorgang", "beschreibung", "beschreibung1", "avisierungstext",
        "transaktionsbeschreibung", "libellé", "libelle", "libellé opération", "libelle operation",
        "libellé de l'opération", "détail de l'écriture", "detail de l'ecriture", "label", "concepto",
        "descripción", "descripcion", "observaciones", "omschrijving", "omschrijving-1", "mededelingen",
        "notes", "text", "extended details",
    ],
}
DEBIT_WORDS = {"af", "debit", "d", "debet", "db", "dr", "s", "soll", "dare", "uscita", "addebito", "cargo", "débit",
               "debito", "dbit", "debito/dare", "a debito"}
CREDIT_WORDS = {"bij", "credit", "c", "cr", "h", "haben", "avere", "entrata", "accredito", "abono", "crédit",
                "credito", "crdt", "a credito"}
PENDING_WORDS = ("pending", "in sospeso", "sospeso", "in attesa", "autorizzat", "non contabilizzat", "vorgemerkt",
                 "offen", "en attente", "pendiente", "declined", "reverted", "failed", "rifiutat", "annullat",
                 "cancel", "storniert", "abgelehnt", "rejected", "refusé", "pdng")
RE_SUMMARY = re.compile(
    r"^(saldo (iniziale|finale|contabile|disponibile|precedente|al \d)|(beginning|ending|opening|closing) balance"
    r"|balance (as of|forward|brought forward|carried forward)|(alter |neuer )?kontostand|anfangssaldo|endsaldo"
    r"|solde (initial|final|au|de)|(totale?|summe|total)\s*:?\s*$)", re.I)
RE_PAYMENT = re.compile(r"payment|thank you|autopay|pagamento ricevuto|pago recibido|zahlung erhalten|"
                        r"remboursement carte|paiement reçu", re.I)
# categorie delle banche (inglese, tedesco, francese, spagnolo) -> parola che il categorizzatore conosce
CATEGORY_HINTS = [
    (("grocer", "supermarket", "lebensmittel", "alimentation", "supermercado", "boodschappen"), "alimentari"),
    (("eating out", "restaurant", "dining", "food & drink", "food and drink", "gastronomie", "restauración",
      "restauracion", "bars"), "ristoranti"),
    (("fuel", "petrol", "gasoline", "tanken", "carburant", "gasolina", "gas station", "gas"), "carburante"),
    (("travel", "holiday", "airline", "reisen", "urlaub", "voyage", "viajes", "hotel"), "viaggi"),
    (("shopping", "shops", "einkauf", "achats", "compras", "merchandise"), "shopping"),
    (("entertainment", "leisure", "freizeit", "loisirs", "ocio"), "svago"),
    (("bills", "utilities", "rechnungen", "factures", "suministros", "energie", "energy"), "bollette"),
    (("health", "pharmacy", "medical", "gesundheit", "santé", "sante", "salud", "farmacia"), "salute"),
    (("cash", "atm", "bargeld", "retrait", "cajero", "geldautomaat"), "prelievo contanti"),
    (("salary", "gehalt", "lohn", "salaire", "nómina", "nomina", "payroll"), "stipendio"),
    (("fees", "fee", "charges", "gebühren", "frais", "comisiones", "fees & adjustments"), "commissioni"),
    (("personal care", "beauty", "körperpflege", "beauté"), "cura personale"),
    (("education", "bildung", "éducation", "educación"), "istruzione"),
    (("gift", "donation", "charity", "spenden", "dons"), "regali"),
    (("insurance", "versicherung", "assurance", "seguros"), "assicurazioni"),
    (("rent", "miete", "loyer", "alquiler", "mortgage", "hypothek"), "affitto"),
    (("home", "household", "haushalt", "maison", "hogar"), "casa"),
    (("subscription", "abonnement", "abo", "suscripci"), "abbonamenti"),
    (("interest", "zinsen", "intérêts", "intereses"), "interessi"),
    (("sport", "fitness"), "sport"),
    (("pets", "haustier", "animaux", "mascotas"), "animali"),
    (("kids", "children", "family", "kinder", "familie", "enfants"), "figli"),
    (("transport", "verkehr", "transporte", "mobilit"), "trasporti"),
    (("tax", "steuer", "impôt", "impuesto"), "tasse"),
]


def _norm(s) -> str:
    s = unicodedata.normalize("NFKC", str(s if s is not None else "")).replace("﻿", "")
    s = s.strip().strip('"').strip("'").lower().replace("_", " ")
    s = re.sub(r"\s+", " ", s).strip()
    return s.rstrip(":").strip()


def _fold(s: str) -> str:
    """Senza accenti (per confronti tolleranti)."""
    return "".join(c for c in unicodedata.normalize("NFKD", s) if not unicodedata.combining(c))


# ---------------------------------------------------------------- profili delle banche
@dataclass
class Profile:
    id: str
    name: str
    country: str          # codice paese ISO (per la bandiera e il gruppo), "EU" = più paesi
    how: str              # come esportare il file (italiano; tradotto con t())
    ccy: str = "EUR"      # valuta se il file non ne indica una
    dayfirst: bool = True  # 03/04/2025 = 3 aprile (False: 4 marzo, stile USA)
    decimal: str | None = None   # "," oppure "."; None = si capisce dai valori
    sign: int = 1          # -1: nel file le spese sono positive (molte carte di credito)
    card: bool = False     # estratto di una carta di credito
    files: str = "CSV, Excel"
    cols: dict[str, list] = field(default_factory=dict)   # nomi propri delle colonne di questa banca
    desc: list[list[str]] = field(default_factory=list)   # colonne da unire per la descrizione, in ordine
    headerless: list[str | None] | None = None             # file senza intestazione: ruolo di ogni colonna
    payee_by_sign: tuple[list[str], list[str]] | None = None  # (colonna per le entrate, colonna per le uscite)
    special: str | None = None   # "fineco" / "revolut" / "chase": lettore dedicato
    note: str = ""

    @property
    def label(self) -> str:
        return f"{flag(self.country)} {self.name}"


FLAGS = {"IT": "🇮🇹", "DE": "🇩🇪", "FR": "🇫🇷", "ES": "🇪🇸", "GB": "🇬🇧", "NL": "🇳🇱", "BE": "🇧🇪", "CH": "🇨🇭",
         "AT": "🇦🇹", "IE": "🇮🇪", "US": "🇺🇸", "EU": "🇪🇺", "LT": "🇱🇹"}


def flag(cc: str) -> str:
    return FLAGS.get(cc, "🏦")


def _p(id, name, country, how, **kw) -> Profile:
    return Profile(id, name, country, how, **kw)


# Nota per i testi «how»: indicano il percorso tipico; le banche a volte rinominano i menu.
PROFILES: list[Profile] = [
    # ------------------------------------------------ Italia
    _p("fineco", "Fineco", "IT", "", special="fineco", files="Excel, CSV"),
    _p("revolut", "Revolut", "EU", "", special="revolut", files="Excel, CSV"),
    _p("intesa", "Intesa Sanpaolo", "IT",
       "Dal sito o dall'app Intesa Sanpaolo apri la lista movimenti del conto (o della carta), scegli il periodo e "
       "usa **Esporta → Excel**. Carica il file `.xlsx` così com'è.",
       files="Excel",
       cols={"date": ["data contabile", "data"], "date2": ["data valuta"], "amount": ["importo"],
             "credit": ["accrediti"], "debit": ["addebiti"], "ccy": ["valuta", "divisa"],
             "status": ["contabilizzazione"], "category": ["categoria"]},
       desc=[["operazione", "descrizione"], ["dettagli", "descrizione estesa"]]),
    _p("unicredit", "UniCredit", "IT",
       "Da Banca via Internet UniCredit apri *Conti → Movimenti*, scegli il periodo e scarica il file in formato "
       "**CSV** (testo con delimitatori). Va bene anche l'**Excel**.",
       cols={"date": ["data", "data registrazione", "data operazione"], "date2": ["valuta", "data valuta"],
             "amount": ["euro", "importo", "importo (eur)"]},
       desc=[["descrizione"]]),
    _p("bancoposta", "BancoPosta (Poste Italiane)", "IT",
       "Dal sito poste.it o dall'app BancoPosta apri la **lista movimenti** del conto, scegli il periodo e "
       "scarica il file **CSV** (o Excel).",
       decimal=",",
       cols={"date": ["data", "data contabile"], "date2": ["valuta", "data valuta"], "credit": ["accrediti"],
             "debit": ["addebiti"]},
       desc=[["descrizione operazioni", "descrizione"]]),
    _p("bper", "BPER Banca", "IT",
       "Dall'internet banking BPER (Smart Web) apri i movimenti del conto, scegli il periodo e esporta in "
       "**CSV** o **Excel**."),
    _p("bancobpm", "Banco BPM", "IT",
       "Dall'internet banking Banco BPM (YouWeb) apri i movimenti del conto, scegli il periodo e esporta in "
       "**Excel** o **CSV**."),
    _p("credem", "Credem", "IT",
       "Dall'internet banking Credem apri i movimenti del conto, scegli il periodo e esporta in **CSV** o **Excel**."),
    _p("mediolanum", "Banca Mediolanum", "IT",
       "Dall'area riservata Banca Mediolanum apri i movimenti del conto, scegli il periodo e esporta in **Excel**."),
    _p("bnl", "BNL BNP Paribas", "IT",
       "Dall'area clienti BNL apri i movimenti del conto, scegli il periodo e scarica in **Excel** o **CSV**."),
    _p("ingit", "ING Italia", "IT",
       "Dall'area riservata ING (Conto Corrente Arancio) apri i movimenti, scegli il periodo e scarica in "
       "**Excel** o **CSV**."),
    _p("credit_agricole_it", "Crédit Agricole Italia", "IT",
       "Dall'internet banking Crédit Agricole Italia apri i movimenti del conto, scegli il periodo e esporta in "
       "**Excel** o **CSV**."),
    _p("sella", "Banca Sella", "IT",
       "Dall'internet banking Banca Sella apri i movimenti del conto, scegli il periodo e esporta in **Excel** o "
       "**CSV**."),
    _p("satispay", "Satispay", "IT",
       "Dall'app Satispay: *Profilo → Esporta transazioni* (ricevi il file **CSV** via email).",
       decimal=",",
       cols={"date": ["date"], "amount": ["amount"], "ccy": ["currency"], "status": ["state"]},
       desc=[["name"], ["extra info"]]),
    # ------------------------------------------------ Europa (più paesi)
    _p("n26", "N26", "EU",
       "Dall'app o dal sito N26: *Conto → Estratti conto* (o *Movimenti*) → **Scarica CSV** e scegli il periodo.",
       decimal=".",
       cols={"date": ["booking date", "date", "datum"], "date2": ["value date"],
             "amount": ["amount (eur)", "betrag (eur)", "montant (eur)", "importo (eur)", "importe (eur)"],
             "ccy": [], "category": ["category", "kategorie", "catégorie"]},
       desc=[["partner name", "payee", "empfänger", "bénéficiaire"],
             ["payment reference", "verwendungszweck", "référence de paiement"]]),
    _p("wise", "Wise", "EU",
       "Dal sito o dall'app Wise: apri il saldo in una valuta → **Estratti conto** → formato **CSV** → scegli il "
       "periodo. Va scaricato un file per ogni valuta.",
       decimal=".",
       cols={"date": ["date", "date time"], "amount": ["amount"], "ccy": ["currency"]},
       desc=[["merchant", "payee name", "payer name"], ["description"], ["payment reference"]]),
    _p("bunq", "bunq", "NL",
       "Dall'app bunq: apri il conto → *Esporta estratto* → formato **CSV** → scegli il periodo.",
       cols={"date": ["datum", "date"], "amount": ["bedrag", "amount"]},
       desc=[["naam", "name"], ["omschrijving", "description"]]),
    # ------------------------------------------------ Germania
    _p("deutschebank", "Deutsche Bank", "DE",
       "Nell'online banking Deutsche Bank apri l'elenco movimenti (*Umsätze*), scegli il periodo e scarica in "
       "**CSV**. Va bene anche **CAMT.053** o **MT940**.",
       files="CSV, CAMT, MT940",
       cols={"date": ["buchungstag", "booking date"], "date2": ["wert", "value date"],
             "amount": ["betrag", "amount"], "debit": ["soll", "debit"], "credit": ["haben", "credit"],
             "ccy": ["währung", "currency"]},
       desc=[["begünstigter / auftraggeber", "beneficiary / originator"], ["verwendungszweck", "payment details"],
             ["umsatzart", "transaction type"]]),
    _p("commerzbank", "Commerzbank", "DE",
       "Nell'online banking Commerzbank apri *Umsätze*, scegli il periodo e usa **Export → CSV**.",
       cols={"date": ["buchungstag"], "date2": ["wertstellung"], "amount": ["betrag"], "ccy": ["währung"],
             "category": ["kategorie"]},
       desc=[["buchungstext"], ["umsatzart"]]),
    _p("sparkasse", "Sparkasse", "DE",
       "Nell'online banking Sparkasse apri *Umsätze*, scegli il periodo e esporta in **CSV-CAMT V2** (o CSV-MT940). "
       "Vanno bene anche i file **CAMT** (anche in ZIP) e **MT940**.",
       files="CSV, CAMT, MT940",
       cols={"date": ["buchungstag", "buchungsdatum"], "date2": ["valutadatum"],
             "amount": ["betrag", "buchungsbetrag"], "ccy": ["waehrung", "währung", "buchungswährung"],
             "status": ["info"]},
       desc=[["beguenstigter/zahlungspflichtiger", "begünstigter/zahlungspflichtiger",
              "auftraggeber/beguenstigter", "auftraggeber/begünstigter"],
             ["verwendungszweck", "transaktionsbeschreibung"], ["buchungstext"]]),
    _p("volksbank", "Volksbank / Raiffeisenbank", "DE",
       "Nell'online banking della tua Volksbank/Raiffeisenbank apri *Umsätze*, scegli il periodo e esporta in "
       "**CSV**. Vanno bene anche **CAMT** e **MT940**.",
       files="CSV, CAMT, MT940",
       cols={"date": ["buchungstag"], "date2": ["valutadatum", "valuta"], "amount": ["betrag", "umsatz"],
             "ccy": ["waehrung", "währung"], "dc": ["soll/haben"], "category": ["kategorie"]},
       desc=[["name zahlungsbeteiligter", "empfänger/zahlungspflichtiger", "auftraggeber/zahlungsempfänger"],
             ["verwendungszweck", "vorgang/verwendungszweck"], ["buchungstext"]]),
    _p("ingde", "ING Deutschland", "DE",
       "Nell'online banking ING apri il conto → *Umsätze* → scegli il periodo → icona **Export (CSV)**.",
       cols={"date": ["buchung", "buchungsdatum"], "date2": ["wertstellungsdatum", "valuta"],
             "amount": ["betrag"], "ccy": ["währung"], "category": ["kategorie"]},
       desc=[["auftraggeber/empfänger"], ["verwendungszweck"], ["buchungstext"]]),
    _p("dkb", "DKB", "DE",
       "Nell'online banking DKB apri il conto → *Umsätze* → scegli il periodo → **Download CSV**.",
       cols={"date": ["buchungsdatum", "buchungstag", "belegdatum"], "date2": ["wertstellung"],
             "amount": ["betrag (€)", "betrag (eur)", "betrag"], "status": ["status"]},
       payee_by_sign=(["zahlungspflichtige*r", "auftraggeber / begünstigter"],
                      ["zahlungsempfänger*in", "auftraggeber / begünstigter"]),
       desc=[["verwendungszweck", "beschreibung"], ["buchungstext", "umsatztyp"]]),
    _p("comdirect", "comdirect", "DE",
       "Nell'online banking comdirect apri *Umsätze* del conto (o della carta), scegli il periodo e usa "
       "**Umsätze exportieren → CSV**.",
       cols={"date": ["buchungstag"], "date2": ["wertstellung (valuta)", "umsatztag"],
             "amount": ["umsatz in eur"]},
       desc=[["buchungstext"], ["vorgang"]]),
    _p("postbank", "Postbank", "DE",
       "Nell'online banking Postbank apri *Umsätze*, scegli il periodo e scarica in **CSV**.",
       cols={"date": ["buchungstag", "buchungsdatum"], "date2": ["wert", "wertstellung"],
             "amount": ["betrag (€)", "betrag (eur)", "betrag"], "debit": ["soll"], "credit": ["haben"]},
       desc=[["begünstigter / auftraggeber", "auftraggeber", "empfänger"], ["verwendungszweck"], ["umsatzart"]]),
    # ------------------------------------------------ Austria / Svizzera
    _p("erste", "Erste Bank / Sparkasse (George)", "AT",
       "In George apri il conto → *Umsätze* → scegli il periodo → **Export → CSV**. Va bene anche **CAMT**.",
       files="CSV, CAMT"),
    _p("ubs", "UBS", "CH",
       "Nell'e-banking UBS apri le transazioni del conto, scegli il periodo e esporta in **CSV**. Va bene anche "
       "**CAMT.053**.", ccy="CHF", files="CSV, CAMT",
       cols={"date": ["abschlussdatum", "trade date", "buchungsdatum", "booking date", "date de transaction"],
             "date2": ["valutadatum", "value date"], "amount": ["einzelbetrag", "individual amount"],
             "debit": ["belastung", "debit", "débit"], "credit": ["gutschrift", "credit", "crédit"],
             "ccy": ["währung", "ccy.", "currency", "monnaie"]},
       desc=[["beschreibung1", "description 1", "description1"], ["beschreibung2", "description 2"],
             ["beschreibung3", "description 3"]]),
    _p("postfinance", "PostFinance", "CH",
       "In e-finance apri *Movimenti*, scegli il periodo e esporta in **CSV**. Va bene anche **CAMT.053**.",
       ccy="CHF", files="CSV, CAMT",
       cols={"date": ["buchungsdatum", "date", "data"], "date2": ["valuta", "valeur"],
             "credit": ["gutschrift in chf", "gutschrift", "crédit en chf", "accredito in chf", "credit in chf"],
             "debit": ["lastschrift in chf", "lastschrift", "débit en chf", "addebito in chf", "debit in chf"]},
       desc=[["avisierungstext", "texte de notification", "testo di avviso", "notification text"]]),
    # ------------------------------------------------ Francia
    _p("bnp", "BNP Paribas", "FR",
       "Dal sito Mes Comptes BNP Paribas apri il conto → **Télécharger** le operazioni → scegli il periodo e il "
       "formato **CSV** o **Excel** (va bene anche **OFX/QIF**).", files="CSV, Excel, OFX, QIF"),
    _p("sg", "Société Générale", "FR",
       "Dal sito Société Générale apri il conto → **Télécharger les opérations** → formato **CSV** (va bene anche "
       "**OFX/QIF**).", files="CSV, OFX, QIF",
       cols={"date": ["date de l'opération", "date de l'operation", "date"],
             "amount": ["montant de l'opération", "montant de l'operation", "montant"], "ccy": ["devise"]},
       desc=[["libellé", "libelle"], ["détail de l'écriture", "detail de l'ecriture"]]),
    _p("ca_fr", "Crédit Agricole", "FR",
       "Dal sito della tua Caisse régionale apri il conto → **Télécharger** → formato **Excel/CSV** (va bene anche "
       "**OFX/QIF**).", files="CSV, Excel, OFX, QIF",
       cols={"date": ["date", "date opération"], "debit": ["débit euros", "debit euros", "débit"],
             "credit": ["crédit euros", "credit euros", "crédit"]},
       desc=[["libellé", "libelle"]]),
    _p("boursorama", "BoursoBank (Boursorama)", "FR",
       "Dal sito BoursoBank apri il conto → **Exporter** le operazioni → formato **CSV**.",
       cols={"date": ["dateop"], "date2": ["dateval"], "amount": ["amount"], "category": ["category"]},
       desc=[["supplierfound"], ["label"], ["comment"]]),
    _p("fortuneo", "Fortuneo", "FR",
       "Dal sito Fortuneo apri il conto → **Exporter** → formato **CSV**.",
       cols={"date": ["date opération", "date operation"], "date2": ["date valeur"],
             "debit": ["débit", "debit"], "credit": ["crédit", "credit"]},
       desc=[["libellé", "libelle"]]),
    # ------------------------------------------------ Spagna
    _p("santander_es", "Banco Santander", "ES",
       "Dalla web Santander apri i movimenti del conto, scegli il periodo e scarica in **Excel**.",
       cols={"date": ["fecha operación", "fecha operacion", "fecha"], "date2": ["fecha valor"],
             "amount": ["importe", "importe eur", "importe (€)"]},
       desc=[["concepto", "descripción", "descripcion"]]),
    _p("bbva", "BBVA", "ES",
       "Dalla web o dall'app BBVA apri i movimenti del conto, scegli il periodo e scarica in **Excel**.",
       cols={"date": ["fecha"], "date2": ["f.valor", "f. valor"], "amount": ["importe"], "ccy": ["divisa"]},
       desc=[["concepto"], ["movimiento"], ["observaciones"]]),
    _p("caixabank", "CaixaBank", "ES",
       "Da CaixaBankNow apri i movimenti del conto, scegli il periodo e scarica in **Excel**.",
       cols={"date": ["fecha", "fecha operación"], "date2": ["fecha valor"], "amount": ["importe", "importe (€)"],
             "ccy": ["divisa"]},
       desc=[["concepto", "concepto común", "concepto propio"], ["más datos", "mas datos"]]),
    _p("inges", "ING España", "ES",
       "Dall'area clienti ING apri i movimenti, scegli il periodo e scarica in **Excel** o **CSV**.",
       cols={"date": ["f. valor", "fecha valor", "fecha"], "amount": ["importe (€)", "importe"],
             "category": ["categoría", "categoria"]},
       desc=[["descripción", "descripcion"], ["comentario"]]),
    # ------------------------------------------------ Regno Unito / Irlanda
    _p("barclays", "Barclays", "GB",
       "Nell'online banking Barclays apri il conto → **Export** → formato **CSV** (va bene anche **OFX/QIF**).",
       ccy="GBP", files="CSV, OFX, QIF",
       cols={"date": ["date"], "amount": ["amount"], "category": ["subcategory"]},
       desc=[["memo"]]),
    _p("hsbc_uk", "HSBC UK", "GB",
       "Nell'online banking HSBC apri il conto → *Download transactions* → formato **CSV** (va bene anche "
       "**OFX/QIF**).", ccy="GBP", files="CSV, OFX, QIF",
       cols={"date": ["date"], "amount": ["amount"]},
       desc=[["description"]], headerless=["date", "desc", "amount"]),
    _p("lloyds", "Lloyds / Halifax / Bank of Scotland", "GB",
       "Nell'internet banking apri il conto → **Export** (Download transactions) → formato **CSV** (va bene anche "
       "**OFX/QIF**).", ccy="GBP", files="CSV, OFX, QIF",
       cols={"date": ["transaction date"], "debit": ["debit amount"], "credit": ["credit amount"]},
       desc=[["transaction description"], ["transaction type"]]),
    _p("natwest", "NatWest / RBS / Ulster Bank", "GB",
       "Nell'online banking apri *Statements → Download or export transactions* → formato **CSV** "
       "(va bene anche **OFX/QIF**).", ccy="GBP", files="CSV, OFX, QIF",
       cols={"date": ["date"], "amount": ["value", "amount"]},
       desc=[["description"], ["type"]]),
    _p("nationwide", "Nationwide", "GB",
       "Nell'internet banking Nationwide apri il conto → **Download transactions** → formato **CSV**.",
       ccy="GBP",
       cols={"date": ["date"], "debit": ["paid out"], "credit": ["paid in"]},
       desc=[["description"], ["transaction type"]]),
    _p("monzo", "Monzo", "GB",
       "Dall'app Monzo: conto → *Statements / Export* → formato **CSV** → scegli il periodo.",
       ccy="GBP", decimal=".",
       cols={"date": ["date"], "amount": ["amount"], "ccy": ["currency"], "category": ["category"]},
       desc=[["name"], ["description"], ["notes and #tags"]]),
    _p("starling", "Starling Bank", "GB",
       "Dall'app Starling: *Account → Statements → Export* → formato **CSV**.",
       ccy="GBP", decimal=".",
       cols={"date": ["date"], "amount": ["amount (gbp)", "amount"], "category": ["spending category"]},
       desc=[["counter party"], ["reference"]]),
    _p("aib", "AIB", "IE",
       "Nell'online banking AIB apri il conto → *Export transactions* → formato **CSV**.",
       cols={"date": ["posted transactions date", "transaction date", "date"],
             "debit": ["debit amount"], "credit": ["credit amount"]},
       desc=[["description1", "description"], ["description2"], ["description3"]]),
    _p("boi", "Bank of Ireland", "IE",
       "In Banking 365 apri il conto → *Export* → formato **CSV**.",
       cols={"date": ["date"], "debit": ["debit"], "credit": ["credit"]}, desc=[["details", "description"]]),
    # ------------------------------------------------ Paesi Bassi / Belgio
    _p("ingnl", "ING Nederland", "NL",
       "Nell'app o in Mijn ING apri il conto → **Af- en bijschrijvingen downloaden** → formato **CSV** (puntkomma "
       "o komma) → scegli il periodo. Va bene anche **MT940/CAMT**.", files="CSV, MT940, CAMT",
       cols={"date": ["datum", "date"], "amount": ["bedrag (eur)", "amount (eur)"],
             "dc": ["af bij", "debit/credit"]},
       desc=[["naam / omschrijving", "name / description"], ["mededelingen", "notifications"]]),
    _p("abnamro", "ABN AMRO", "NL",
       "In Internet Bankieren ABN AMRO apri *Mutaties downloaden* → formato **TXT** (tab-gescheiden) o **CSV**; "
       "vanno bene anche **MT940/CAMT**.", files="TXT, CSV, MT940, CAMT", decimal=",",
       headerless=["account", "ccy", "date", None, None, "date2", "amount", "desc"]),
    _p("rabobank", "Rabobank", "NL",
       "In Rabo Online Bankieren apri il conto → **Exporteren** → formato **CSV** → scegli il periodo.",
       decimal=",",
       cols={"date": ["datum"], "date2": ["rentedatum"], "amount": ["bedrag"], "ccy": ["munt"]},
       desc=[["naam tegenpartij"], ["omschrijving-1"], ["omschrijving-2"], ["omschrijving-3"]]),
    _p("kbc", "KBC", "BE",
       "In KBC Touch apri il conto → **Export** → formato **CSV**.", decimal=",",
       cols={"date": ["datum", "date"], "date2": ["valuta"], "amount": ["bedrag", "montant"], "ccy": ["munt"]},
       desc=[["naam tegenpartij"], ["omschrijving", "vrije mededeling"]]),
    _p("bnpfortis", "BNP Paribas Fortis", "BE",
       "In Easy Banking Web apri il conto → **Exporteren / Exporter** → formato **CSV**.", decimal=",",
       cols={"date": ["uitvoeringsdatum", "date d'exécution", "execution date"],
             "date2": ["valutadatum", "date valeur", "value date"], "amount": ["bedrag", "montant", "amount"],
             "ccy": ["valuta rekening", "devise du compte", "account currency"]},
       desc=[["naam van de tegenpartij", "nom de la contrepartie", "tegenpartij van de verrichting"],
             ["mededeling", "communication"], ["details", "détails"]]),
    # ------------------------------------------------ Stati Uniti
    _p("chase", "Chase", "US", "", special="chase", ccy="USD", dayfirst=False, files="CSV, QFX/OFX"),
    _p("bofa", "Bank of America", "US",
       "Su bankofamerica.com apri il conto o la carta → **Download** (sopra l'elenco movimenti) → formato "
       "**Microsoft Excel (CSV)** o **Quicken (QFX)** → scegli il periodo.",
       ccy="USD", dayfirst=False, decimal=".", files="CSV, QFX/OFX",
       cols={"date": ["date", "posted date"], "amount": ["amount"]},
       desc=[["payee", "description"]]),
    _p("wellsfargo", "Wells Fargo", "US",
       "Su wellsfargo.com: *Account activity → Download account activity* → formato **Comma delimited (CSV)** o "
       "**Quicken (QFX)**.", ccy="USD", dayfirst=False, decimal=".", files="CSV, QFX/OFX",
       headerless=["date", "amount", None, None, "desc"]),
    _p("citi", "Citi", "US",
       "Su citi.com apri il conto o la carta → **Download transactions** → formato **CSV** (o QFX).",
       ccy="USD", dayfirst=False, decimal=".", files="CSV, QFX/OFX",
       cols={"date": ["date", "transaction date"], "debit": ["debit"], "credit": ["credit"], "amount": ["amount"],
             "status": ["status"]},
       desc=[["description"]]),
    _p("capitalone", "Capital One", "US",
       "Su capitalone.com apri il conto o la carta → **Download transactions** → formato **CSV** (o QFX).",
       ccy="USD", dayfirst=False, decimal=".", files="CSV, QFX/OFX",
       cols={"date": ["transaction date", "date"], "date2": ["posted date"], "debit": ["debit"],
             "credit": ["credit"], "amount": ["transaction amount", "amount"],
             "dc": ["transaction type"], "category": ["category"]},
       desc=[["description", "transaction description"]]),
    _p("amex", "American Express", "US",
       "Sul sito American Express apri l'estratto → **Download** → formato **CSV** (o Excel / QFX). "
       "Nei file Amex gli acquisti sono positivi: il segno viene corretto da solo.",
       ccy="USD", dayfirst=False, decimal=".", sign=-1, card=True, files="CSV, Excel, QFX/OFX",
       cols={"date": ["date"], "amount": ["amount"], "category": ["category"]},
       desc=[["description"], ["extended details"]]),
    _p("discover", "Discover", "US",
       "Su discover.com: *Activity & Statements* → **Download** → formato **CSV** (o QFX). Nei file Discover gli "
       "acquisti sono positivi: il segno viene corretto da solo.",
       ccy="USD", dayfirst=False, decimal=".", sign=-1, card=True, files="CSV, QFX/OFX",
       cols={"date": ["trans. date", "transaction date"], "date2": ["post date"], "amount": ["amount"],
             "category": ["category"]},
       desc=[["description"]]),
    _p("usbank", "U.S. Bank", "US",
       "Su usbank.com apri il conto → **Download transactions** → formato **CSV** (o QFX).",
       ccy="USD", dayfirst=False, decimal=".", files="CSV, QFX/OFX",
       cols={"date": ["date"], "amount": ["amount"]},
       desc=[["name"], ["memo"]]),
    _p("ally", "Ally Bank", "US",
       "Su ally.com apri il conto → **Download** → formato **CSV** (o QFX).",
       ccy="USD", dayfirst=False, decimal=".", files="CSV, QFX/OFX",
       cols={"date": ["date"], "amount": ["amount"]}, desc=[["description"]]),
    _p("schwab", "Charles Schwab Bank", "US",
       "Su schwab.com apri il conto → *History* → **Export** → formato **CSV**.",
       ccy="USD", dayfirst=False, decimal=".", files="CSV, QFX/OFX",
       cols={"date": ["date"], "debit": ["withdrawal"], "credit": ["deposit"], "status": ["status"]},
       desc=[["description"]]),
    # ------------------------------------------------ qualsiasi banca
    _p("auto", "Altra banca: riconoscimento automatico", "",
       "Carica il file della tua banca (CSV, Excel, OFX/QFX, QIF, CAMT o MT940): l'app cerca da sola le colonne "
       "di data, descrizione e importo. Se non ci riesce, puoi indicarle tu.",
       files="CSV, Excel, OFX/QFX, QIF, CAMT, MT940"),
]
BY_ID = {p.id: p for p in PROFILES}
COUNTRY_ORDER = ["IT", "EU", "LT", "DE", "AT", "CH", "FR", "ES", "GB", "IE", "NL", "BE", "US", ""]
UPLOAD_TYPES = ["csv", "txt", "tab", "xlsx", "xls", "xlsm", "ofx", "qfx", "qif", "xml", "sta", "mt940", "940", "zip"]


def ordered() -> list[Profile]:
    """Fineco, Revolut e Chase in testa (già usate), poi per paese, «Altra banca» in fondo."""
    first = [BY_ID["fineco"], BY_ID["revolut"], BY_ID["chase"]]
    rest = [p for p in PROFILES if p not in first and p.id != "auto"]
    rest.sort(key=lambda p: (COUNTRY_ORDER.index(p.country) if p.country in COUNTRY_ORDER else 99, p.name.lower()))
    return first + rest + [BY_ID["auto"]]


# ---------------------------------------------------------------- lettura del file in righe
def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1", errors="replace")


def _csv_rows(text: str) -> list[list]:
    lines = text.splitlines()
    sample = "\n".join(lines[:120])
    best, best_score = ",", -1
    for d in (";", ",", "\t", "|"):
        try:
            counts = [len(r) for r in csv.reader(io.StringIO(sample), delimiter=d) if any(c.strip() for c in r)]
        except csv.Error:
            continue
        counts = [c for c in counts if c > 1]
        if not counts:
            continue
        modal = max(set(counts), key=counts.count)
        score = counts.count(modal) * modal
        if score > best_score:
            best, best_score = d, score
    reader = csv.reader(io.StringIO(text), delimiter=best)
    return [r for r in reader]


def read_tables(data: bytes, filename: str) -> list[tuple[str, list[list]]]:
    """[(nome del foglio, righe)] per CSV/TXT ed Excel (tutti i fogli)."""
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")) or data[:4] == b"PK\x03\x04":
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        out = [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]
        wb.close()
        return out
    if name.endswith(".xls") or data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        try:
            import xlrd
            book = xlrd.open_workbook(file_contents=data)
            out = []
            for sh in book.sheets():
                rows = []
                for i in range(sh.nrows):
                    row = []
                    for c in sh.row(i):
                        if c.ctype == xlrd.XL_CELL_DATE:
                            try:
                                row.append(xlrd.xldate_as_datetime(c.value, book.datemode))
                            except Exception:
                                row.append(c.value)
                        else:
                            row.append(c.value)
                    rows.append(row)
                out.append((sh.name, rows))
            return out
        except Exception:
            pass  # alcune banche salvano come .xls una pagina HTML o un CSV
        head = data[:2000].lower()
        if b"<html" in head or b"<table" in head:
            tables = pd.read_html(io.BytesIO(data), header=None)
            return [(f"tab{i + 1}", t.astype(object).where(t.notna(), None).values.tolist())
                    for i, t in enumerate(tables)]
    return [("", _csv_rows(_decode(data)))]


# ---------------------------------------------------------------- date e importi
MONTHS = {
    # italiano, inglese, tedesco, francese, spagnolo, olandese (prime lettere)
    "gen": 1, "jan": 1, "ene": 1, "janv": 1, "januar": 1, "feb": 2, "fév": 2, "fev": 2, "févr": 2, "mar": 3,
    "mär": 3, "maer": 3, "mrt": 3, "apr": 4, "avr": 4, "abr": 4, "mag": 5, "may": 5, "mai": 5, "mei": 5,
    "giu": 6, "jun": 6, "juin": 6, "lug": 7, "jul": 7, "juil": 7, "ago": 8, "aug": 8, "aoû": 8, "aou": 8,
    "set": 9, "sep": 9, "sept": 9, "ott": 10, "oct": 10, "okt": 10, "nov": 11, "dic": 12, "dec": 12, "dez": 12,
    "déc": 12, "dic.": 12,
}


def _month(word: str) -> int | None:
    w = word.lower().strip(".")
    for n in (5, 4, 3):
        if w[:n] in MONTHS:
            return MONTHS[w[:n]]
    return None


def _mk(y: int, m: int, d: int) -> str | None:
    if y < 100:
        y += 2000 if y < 70 else 1900
    try:
        return date(y, m, d).isoformat()
    except ValueError:
        return None


RE_ISO = re.compile(r"^(\d{4})[-/.](\d{1,2})[-/.](\d{1,2})")
RE_COMPACT = re.compile(r"^(\d{4})(\d{2})(\d{2})$")
RE_NUM = re.compile(r"^(\d{1,2})([./\-])(\d{1,2})[./\-](\d{2,4})(?!\d)")
RE_TXT1 = re.compile(r"^(\d{1,2})[\s\-.]+([^\W\d_]{3,})\.?[\s\-.,]+(\d{2,4})")
RE_TXT2 = re.compile(r"^([^\W\d_]{3,})\.?\s+(\d{1,2}),?\s+(\d{4})")


def date_parser(values: list, dayfirst: bool):
    """Funzione che converte un valore in data ISO, con l'ordine giorno/mese dedotto dai valori della colonna
    (se qualche data ha il primo numero > 12 è giorno/mese, se il secondo > 12 è mese/giorno)."""
    order = None
    for v in values:
        if isinstance(v, str):
            m = RE_NUM.match(v.strip())
            if m and m.group(2) != ".":
                a, b = int(m.group(1)), int(m.group(3))
                if a > 12 >= b:
                    order = True
                    break
                if b > 12 >= a:
                    order = False
                    break
    df = dayfirst if order is None else order

    def conv(v) -> str | None:
        if v is None:
            return None
        if isinstance(v, (datetime, pd.Timestamp)):
            return None if pd.isna(v) else v.date().isoformat()
        if isinstance(v, date):
            return v.isoformat()
        if isinstance(v, (int, float)):
            if isinstance(v, float) and pd.isna(v):
                return None
            if 20000 < v < 80000:  # data Excel come numero
                return (date(1899, 12, 30) + timedelta(days=int(v))).isoformat()
            s = str(int(v)) if float(v).is_integer() else str(v)
        else:
            s = str(v).strip()
        if not s:
            return None
        m = RE_ISO.match(s)
        if m:
            return _mk(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        m = RE_COMPACT.match(s)
        if m:
            return _mk(int(m.group(1)), int(m.group(2)), int(m.group(3)))
        m = RE_NUM.match(s)
        if m:
            a, sep, b, y = int(m.group(1)), m.group(2), int(m.group(3)), int(m.group(4))
            if sep == "." or df:
                return _mk(y, b, a)
            return _mk(y, a, b)
        m = RE_TXT1.match(s)
        if m and _month(m.group(2)):
            return _mk(int(m.group(3)), _month(m.group(2)), int(m.group(1)))
        m = RE_TXT2.match(s)
        if m and _month(m.group(1)):
            return _mk(int(m.group(3)), _month(m.group(1)), int(m.group(2)))
        return None
    return conv


RE_CCY_TXT = re.compile(r"(eur|usd|gbp|chf|sek|nok|dkk|pln|czk|huf|cad|aud|jpy|fr\.|€|\$|£)", re.I)


def _num_text(v) -> str:
    s = str(v).strip()
    s = s.replace("−", "-").replace("–", "-").replace("\xa0", "").replace(" ", "")
    s = RE_CCY_TXT.sub("", s).replace(" ", "").replace("'", "").replace("’", "")
    return s


def amount_parser(values: list, decimal: str | None):
    """Funzione che converte un importo (testo o numero) in float. Il separatore decimale si deduce dai valori
    della colonna quando non è indicato (1.234,56 / 1,234.56 / 12,5 / 12.50)."""
    votes = {",": 0, ".": 0}
    for v in values:
        if isinstance(v, str):
            s = _num_text(v)
            if re.search(r"\d\.\d{3},\d", s) or re.search(r",\d{1,2}(?!\d)\D*$", s):
                votes[","] += 1
            elif re.search(r"\d,\d{3}\.\d", s) or re.search(r"\.\d{1,2}(?!\d)\D*$", s):
                votes["."] += 1
    col_dec = decimal or ("," if votes[","] > votes["."] else "." if votes["."] > votes[","] else None)

    def conv(v) -> float | None:
        if v is None:
            return None
        if isinstance(v, bool):
            return None
        if isinstance(v, (int, float)):
            return None if pd.isna(v) else float(v)
        s = _num_text(v)
        if s in ("", "-", "+", "--", "nan", "none", "null"):
            return None
        neg = False
        if s.startswith("(") and s.endswith(")"):
            neg, s = True, s[1:-1]
        up = s.upper()
        if up.endswith("CR"):
            s = s[:-2]
        elif up.endswith("DR"):
            neg, s = not neg, s[:-2]
        if s.endswith("-"):           # 12,50- (stile tedesco)
            neg, s = not neg, s[:-1]
        if s.startswith("+"):
            s = s[1:]
        if s.startswith("-"):
            neg, s = not neg, s[1:]
        if not re.fullmatch(r"[\d.,]+", s or "x"):
            return None
        if "," in s and "." in s:          # entrambi: l'ultimo è il decimale
            dec = "," if s.rfind(",") > s.rfind(".") else "."
        elif "," in s:
            dec = "," if not re.fullmatch(r"\d{1,3}(,\d{3})+", s) or col_dec == "," else "."
            if s.count(",") > 1:
                dec = "."
        elif "." in s:
            dec = "." if not re.fullmatch(r"\d{1,3}(\.\d{3})+", s) or col_dec == "." else ","
            if s.count(".") > 1:
                dec = ","
        else:
            dec = "."
        if dec == ",":
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
        try:
            x = float(s)
        except ValueError:
            return None
        return -x if neg else x
    return conv


# ---------------------------------------------------------------- riconoscimento delle colonne
def _cands(profile: Profile, role: str) -> list[str]:
    own = [_norm(c) for c in profile.cols.get(role, [])]
    if role == "ccy" and "ccy" in profile.cols and not profile.cols["ccy"]:
        return []  # la banca non ha una colonna valuta (es. N26: gli importi sono sempre in euro)
    return own + [c for c in (_norm(x) for x in AUTO.get(role, [])) if c not in own]


def _match(header: list[str], cands: list[str], used: set[int]) -> int | None:
    folded = [_fold(h) for h in header]
    for c in cands:  # prima i nomi esatti, nell'ordine di preferenza
        if c.endswith("*"):
            continue
        fc = _fold(c)
        for i, h in enumerate(header):
            if i not in used and (h == c or folded[i] == fc):
                return i
    for c in cands:  # poi i prefissi («betrag (*»)
        if c.endswith("*"):
            p = _fold(c[:-1])
            for i, h in enumerate(folded):
                if i not in used and h.startswith(p):
                    return i
    return None


def _map_header(header: list[str], profile: Profile) -> dict:
    used: set[int] = set()
    cols: dict = {}
    for role in ("date", "amount", "debit", "credit", "dc", "ccy", "status", "category", "date2"):
        i = _match(header, _cands(profile, role), used)
        if i is not None:
            cols[role] = i
            used.add(i)
    if "ccy" in cols and "amount" in cols:  # più colonne «Währung» (ING): quella subito dopo l'importo
        same = [j for j, h in enumerate(header) if h == header[cols["ccy"]]]
        after = [j for j in same if j > cols["amount"]]
        if len(same) > 1 and after:
            cols["ccy"] = after[0]
    # descrizione: colonne indicate dalla banca, altrimenti controparte + causale
    parts: list[int] = []
    groups = profile.desc or []
    for g in groups:
        i = _match(header, [_norm(x) for x in g], used | set(parts))
        if i is not None:
            parts.append(i)
    if not parts:
        i = _match(header, [_norm(x) for x in AUTO["payee"]], used)
        if i is not None:
            parts.append(i)
        for _ in range(2):
            j = _match(header, [_norm(x) for x in AUTO["purpose"]], used | set(parts))
            if j is not None:
                parts.append(j)
    cols["desc"] = parts
    payee = _match(header, [_norm(x) for x in AUTO["payee"]], used)
    if payee is not None:
        cols["payee"] = payee
    if profile.payee_by_sign:
        cin = _match(header, [_norm(x) for x in profile.payee_by_sign[0]], used)
        cout = _match(header, [_norm(x) for x in profile.payee_by_sign[1]], used | ({cin} if cin is not None else set()))
        if cin is not None or cout is not None:
            cols["payee_in"], cols["payee_out"] = cin, cout
    return cols


def _find_header(rows: list[list], profile: Profile) -> tuple[int, dict] | None:
    for i, row in enumerate(rows[:80]):
        header = [_norm(c) for c in row]
        if sum(1 for h in header if h) < 2:
            continue
        cols = _map_header(header, profile)
        if "date" in cols and ("amount" in cols or "debit" in cols or "credit" in cols):
            return i, cols
    return None


# ---------------------------------------------------------------- lettura con un profilo
class FormatError(ValueError):
    pass


@dataclass
class Result:
    rows: list[dict]
    fmt: str                 # "CSV", "Excel", "OFX/QFX", …
    columns: dict[str, str] = field(default_factory=dict)   # ruolo -> nome della colonna usata
    sheet: str = ""
    skipped_pending: int = 0
    skipped_invalid: int = 0
    auto: bool = False       # colonne trovate solo con i nomi generici (il file è diverso dal solito)
    card: bool = False


def _cell(row: list, i: int | None):
    if i is None or i >= len(row):
        return None
    v = row[i]
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def _text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", str(v)).strip()


def _cat_hint(cat: str) -> str:
    c = cat.lower()
    for words, hint in CATEGORY_HINTS:
        if any(w in c for w in words):
            return hint
    return c  # categorie già in italiano (Intesa, Satispay…) le capisce il categorizzatore


def _uses_own_names(header: list[str], cols: dict, profile: Profile) -> bool:
    if not profile.cols:
        return True
    own = {_norm(x) for v in profile.cols.values() for x in v}
    return any(header[cols[r]] in own for r in ("date", "amount", "debit", "credit") if r in cols)


def parse_table(rows: list[list], profile: Profile, fmt: str, sheet: str = "", base_ccy: str = "EUR") -> Result:
    found = _find_header(rows, profile)
    headerless = False
    if found is None:
        if not profile.headerless:
            raise FormatError(i18n.t("non trovo le colonne di data e importo"))
        roles = profile.headerless
        cols = {"desc": []}
        for i, r in enumerate(roles):
            if r == "desc":
                cols["desc"].append(i)
            elif r:
                cols[r] = i
        start, header = 0, [f"#{i + 1}" for i in range(max(len(r) for r in rows) if rows else 0)]
        headerless = True
    else:
        h, cols = found
        header = [_norm(c) for c in rows[h]]
        start = h + 1
    data = [r for r in rows[start:] if any(_cell(r, i) is not None for i in range(len(r)))]
    # descrizioni con il separatore non protetto (UniCredit): campi in più uniti nella descrizione
    if not headerless and cols["desc"] and fmt == "CSV":
        n = len(header)
        fixed = []
        di = cols["desc"][0]
        for r in data:
            r = list(r)
            while len(r) > n and (r[-1] is None or str(r[-1]).strip() == ""):
                r.pop()
            if len(r) > n and di < n:
                extra = len(r) - n
                merged = ";".join(str(x) for x in r[di:di + extra + 1] if x is not None)
                r = r[:di] + [merged] + r[di + extra + 1:]
            fixed.append(r)
        data = fixed
    col_vals = lambda i: [_cell(r, i) for r in data] if i is not None else []  # noqa: E731
    dconv = date_parser(col_vals(cols.get("date")) + col_vals(cols.get("date2")), profile.dayfirst)
    aconv = {k: amount_parser(col_vals(cols.get(k)), profile.decimal) for k in ("amount", "debit", "credit")}
    # colonna valuta: valida solo se contiene codici di 3 lettere
    if "ccy" in cols:
        vals = [str(v).strip().upper() for v in col_vals(cols["ccy"]) if v is not None]
        if not vals or sum(bool(re.fullmatch(r"[A-Z]{3}", v)) for v in vals) < 0.8 * len(vals):
            cols.pop("ccy")
    # colonna senza nome subito dopo l'importo con S/H (vecchi file Volksbank): è il dare/avere
    if "dc" not in cols and "amount" in cols and not headerless:
        j = cols["amount"] + 1
        if j < len(header) and not header[j]:
            cols["dc"] = j
    # colonna dare/avere: valida solo se contiene parole note
    if "dc" in cols:
        vals = [_norm(v) for v in col_vals(cols["dc"]) if v is not None]
        known = DEBIT_WORDS | CREDIT_WORDS
        if not vals or sum(v in known for v in vals) < 0.8 * len(vals):
            cols.pop("dc")
    out: list[dict] = []
    pending = invalid = 0
    default_ccy = profile.ccy if profile.id != "auto" else base_ccy
    for r in data:
        if "status" in cols:
            stt = _norm(_cell(r, cols["status"]) or "")
            if stt and any(w in stt for w in PENDING_WORDS):
                pending += 1
                continue
        d = dconv(_cell(r, cols.get("date"))) or dconv(_cell(r, cols.get("date2")))
        a = aconv["amount"](_cell(r, cols.get("amount"))) if "amount" in cols else None
        if a is None and ("debit" in cols or "credit" in cols):
            de = aconv["debit"](_cell(r, cols.get("debit"))) if "debit" in cols else None
            cr = aconv["credit"](_cell(r, cols.get("credit"))) if "credit" in cols else None
            if de is not None or cr is not None:
                a = abs(cr or 0.0) - abs(de or 0.0)
        if "dc" in cols and a is not None:
            flag_ = _norm(_cell(r, cols["dc"]) or "")
            if flag_ in DEBIT_WORDS:
                a = -abs(a)
            elif flag_ in CREDIT_WORDS:
                a = abs(a)
        parts = [_text(_cell(r, i)) for i in cols["desc"]]
        if "payee_in" in cols and a is not None:
            pi = cols["payee_in"] if a > 0 else cols["payee_out"]
            p = _text(_cell(r, pi))
            if p:
                parts = [p] + parts
        parts = [p for p in dict.fromkeys(parts) if p and p.lower() not in ("nan", "none")]
        desc = " · ".join(parts)
        texts = [str(c).strip() for c in r if isinstance(c, str)]
        if RE_SUMMARY.match(desc or "") or any(RE_SUMMARY.match(x) for x in texts[:2]):
            continue  # righe di saldo o totale
        if _norm(_cell(r, cols.get("date")) or "") in ("offen", "pending", "vorgemerkt", "in sospeso"):
            pending += 1  # movimento non ancora contabilizzato (comdirect scrive «offen» al posto della data)
            continue
        if not d or a is None:
            if any(_cell(r, i) is not None for i in range(len(r))):
                invalid += 1
            continue
        if a == 0:
            continue
        a = round(a * profile.sign, 2)
        ccy = str(_cell(r, cols["ccy"]) or "").strip().upper() if "ccy" in cols else ""
        cat = _text(_cell(r, cols.get("category"))) if "category" in cols else ""
        hint = _cat_hint(cat) if cat else ""
        if profile.card and a > 0 and RE_PAYMENT.search(desc):
            hint = "giroconto"  # pagamento della carta dal conto: non è un'entrata
        payee = _text(_cell(r, cols.get("payee_in") if a > 0 and "payee_in" in cols else
                            cols.get("payee_out") if "payee_out" in cols else cols.get("payee")))
        out.append({"date": d, "description": desc or "-", "amount": float(a), "currency": ccy or default_ccy,
                    "_bank_category": hint, "_counterparty": payee or None})
    if not out:
        raise FormatError(i18n.t("nessun movimento valido trovato"))
    names = {"date": i18n.t("Data"), "date2": i18n.t("Data (riserva)"), "amount": i18n.t("Importo"),
             "debit": i18n.t("Uscite"), "credit": i18n.t("Entrate"), "dc": i18n.t("Dare/avere"),
             "ccy": i18n.t("Valuta"), "category": i18n.t("Categoria"), "status": i18n.t("Stato")}
    raw_header = rows[start - 1] if not headerless else header
    show = {names[k]: _text(raw_header[v]) if not headerless else f"#{v + 1}"
            for k, v in cols.items() if k in names and v is not None}
    if cols["desc"]:
        show[i18n.t("Descrizione")] = " + ".join(_text(raw_header[i]) if not headerless else f"#{i + 1}"
                                                  for i in cols["desc"])
    return Result(out, fmt, show, sheet, pending, invalid,
                  auto=not headerless and not _uses_own_names(header, cols, profile), card=profile.card)


def parse(profile: Profile, data: bytes, filename: str, base_ccy: str = "EUR") -> Result:
    """Legge il file con il profilo della banca. Prova prima i formati standard (OFX/QFX, QIF, CAMT, MT940),
    poi le tabelle (CSV/Excel, tutti i fogli). FormatError se non trova movimenti."""
    from core import stdformats

    kind = stdformats.detect(data, filename)
    if kind:
        fmt, rows = stdformats.parse(data, filename, dayfirst=profile.dayfirst)
        default = profile.ccy if profile.id != "auto" else base_ccy
        for r in rows:  # nei formati standard il segno è già quello giusto (uscite negative)
            r["currency"] = (r.get("currency") or default).upper()
            r["amount"] = round(r["amount"], 2)
        if not rows:
            raise FormatError(i18n.t("nessun movimento valido trovato"))
        return Result(rows, fmt, card=profile.card)
    name = filename.lower()
    fmt = "Excel" if name.endswith((".xlsx", ".xlsm", ".xls")) else "CSV"
    errors = []
    for sheet, rows in read_tables(data, filename):
        if not rows:
            continue
        try:
            return parse_table(rows, profile, fmt, sheet, base_ccy)
        except FormatError as e:
            errors.append(str(e))
    raise FormatError(errors[0] if errors else i18n.t("file vuoto"))
